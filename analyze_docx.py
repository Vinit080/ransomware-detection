import docx
import sys

# Fix printing unicode in Windows
sys.stdout.reconfigure(encoding='utf-8')

def main():
    doc_path = r"E:\Ransomware Det\Ransomware_GenAI_TIFS_28718_2026_Reformatted.docx"
    doc = docx.Document(doc_path)
    
    print("--- Searching for 'future .' ---")
    for i, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                if "future ." in cell.text:
                    print(f"Found 'future .' in Table {i}, Row {r_idx}, Col {c_idx}")
                    print(f"Original Text: {cell.text.encode('ascii', 'ignore').decode()}")
                    new_text = cell.text.replace("future .", "future work.")
                    cell.text = new_text

    # Also search paragraphs just in case
    for p_idx, para in enumerate(doc.paragraphs):
        if "future ." in para.text:
            print(f"Found 'future .' in paragraph {p_idx}")
            para.text = para.text.replace("future .", "future work.")

    print("\n--- Searching for Figure 1 ---")
    for p_idx, para in enumerate(doc.paragraphs):
        text = para.text
        if "Fig. 1" in text or "Figure 1" in text:
            print(f"Found in paragraph {p_idx}: {text.encode('ascii', 'ignore').decode()}")

    print("\n--- Updating Table III metrics ---")
    for i, table in enumerate(doc.tables):
        is_table_iii = False
        for row in table.rows:
            if "ATT&CK mapping accuracy" in row.cells[0].text or "Telemetry tamper detection" in row.cells[0].text:
                is_table_iii = True
                break
        
        if is_table_iii:
            print(f"Found Table III at index {i}")
            for row in table.rows:
                metric = row.cells[0].text.strip()
                if metric == "ATT&CK mapping accuracy":
                    print(f"Updating {metric}: {row.cells[2].text} -> 0.51")
                    row.cells[2].text = "0.51"
                elif metric == "Telemetry tamper detection":
                    print(f"Updating {metric}: {row.cells[2].text} -> 0.83")
                    row.cells[2].text = "0.83"
                elif metric == "CPU overhead":
                    print(f"Updating {metric}: {row.cells[2].text} -> 0.28")
                    row.cells[2].text = "0.28"
    
    modified_path = r"E:\Ransomware Det\Ransomware_GenAI_TIFS_28718_2026_Reformatted_Fixed.docx"
    doc.save(modified_path)
    print(f"\nSaved modified document to {modified_path}")

if __name__ == '__main__':
    main()
